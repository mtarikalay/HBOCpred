from pathlib import Path
from io import BytesIO
from copy import deepcopy
from lxml import etree as E
from docx import Document
from docx.shared import Pt,Inches,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pandas as pd,json
from response_content import INTRO,ITEMS,ADDITIONAL
from reviewer_verbatim import OPENING,RECOMMENDATION
from build_manuscript import table,NS,ROOT
BASE=Path(__file__).resolve().parent;OUT=ROOT/'documentation'

def clear(d):
 for x in list(d._element.body):
  if x.tag!=qn('w:sectPr'):d._element.body.remove(x)

def save_document(d,path):
 buffer=BytesIO();d.save(buffer)
 temp=path.with_suffix('.docx.tmp');temp.write_bytes(buffer.getvalue());temp.replace(path)

def pp(d,text,style='Normal',bold_prefix=None):
 p=d.add_paragraph(style=style)
 if bold_prefix and text.startswith(bold_prefix):p.add_run(bold_prefix).bold=True;p.add_run(text[len(bold_prefix):])
 else:p.add_run(text)
 p.paragraph_format.widow_control=True
 return p
# Preserve the supplied response template rather than changing the manuscript design.
d=Document(BASE/'source_response.docx');clear(d)
pp(d,'Response to Reviewer 1 — Second-Round Report','Title')
pp(d,'Manuscript ID: ijms-4511203')
pp(d,'HBOCpred: an ensemble framework for prioritisation of germline missense variants in hereditary breast and ovarian cancer')
pp(d,'Reviewer opening assessment','Heading 1');pp(d,OPENING)
pp(d,'General response','Heading 1');pp(d,INTRO)
for title,comment,replies,where in ITEMS:
 pp(d,title,'Heading 1')
 pp(d,'Reviewer comment: '+comment,bold_prefix='Reviewer comment: ')
 for i,t in enumerate(replies):pp(d,('Response: ' if i==0 else '')+t,bold_prefix='Response: ' if i==0 else None)
 p=pp(d,'Location and evidence: '+where,bold_prefix='Location and evidence: ');p.paragraph_format.space_after=Pt(9)
pp(d,'Additional corrections identified during the revision','Heading 1')
for t in ADDITIONAL:pp(d,t)
pp(d,'Reviewer recommendation','Heading 1');pp(d,RECOMMENDATION)
pp(d,'Response: The revised paper is limited to internal assessment of computational agreement and research prioritisation. It documents the remaining source-label dependence and uncertainty in application. Independent clinical validation remains necessary, and persistent deposition must precede resubmission.',bold_prefix='Response: ')
save_document(d,OUT/'HBOCpred_Response_Reviewer1_Round2.docx')
# Supplement inherits the IJMS template, page geometry, headers, typography and table styling.
d=Document(OUT/'HBOCpred_IJMS_Reviewer1_Round2_clean.docx');source_table=deepcopy(d.tables[0]._tbl);clear(d)
body='MDPI_3.1_text';head='MDPI_2.2_heading2';cap='MDPI_4.1_table_caption'
pp(d,'Supplementary Materials','MDPI_1.2_title');pp(d,'HBOCpred — audited reanalysis for Reviewer 1, second round',body)
pp(d,'The computational release is labelled R4 to distinguish it from the supplied R2 manuscript and Claude R3 files. All values below derive from the actual R4 fitted objects and saved out-of-fold predictions. The primary cohort contains 2,160 historical source-labelled anchors; 114 do not meet output completeness in each repeated-CV assessment. No external clinical validation is claimed.',body)
def addtable(rows,widths):
 t=deepcopy(source_table);table(t,rows,widths,False)
 pr=t.find('w:tblPr',NS);i=pr.find('w:tblInd',NS)
 if i is not None:i.set(qn('w:w'),'0')
 jc=pr.find('w:jc',NS)
 if jc is not None:jc.set(qn('w:val'),'center')
 d._element.body.insert(len(d._element.body)-1,t)
 return t
pp(d,'S1. Gene-held-out denominators and performance',head)
pp(d,'Table S1. Source and evaluable counts, with conditional class-specific results.',cap)
g=pd.read_csv(ROOT/'tables/per_gene_held_out.csv')
rows=[['Gene','Source n','Scored n','LP/P: correct/scored','Sensitivity (95% interval)','B/LB: correct/scored','Specificity (95% interval)']]
for _,r in g.iterrows():rows.append([r.Gene,int(r.n_total),int(r.n_evaluable),f'{int(r.TP)}/{int(r.TP+r.FN)}',f'{100*r.sensitivity:.1f}% ({100*r.sensitivity_lo:.1f}–{100*r.sensitivity_hi:.1f})',f'{int(r.TN)}/{int(r.TN+r.FP)}',f'{100*r.specificity:.1f}% ({100*r.specificity_lo:.1f}–{100*r.specificity_hi:.1f})'])
addtable(rows,[1100,850,850,1700,2130,1700,2130])
pp(d,'Intervals are Wilson intervals conditional on the evaluated source records. Incomplete counts equal source n minus scored n. These denominators reveal, for example, that PTEN specificity is based on five evaluable B/LB anchors, while 17 further B/LB anchors are incomplete. The full machine-readable table also includes per-gene AUROC, AUPRC, Brier score, confusion counts and V-state counts.',body)
pp(d,'S2. Conditional intervals for the fixed first repeat',head)
pp(d,'A class-stratified bootstrap used one fixed out-of-fold prediction per evaluable variant, 2,000 resamples and seed 20260905. The intervals below condition on the fitted scores and the sampled source-label classes. Models were not refitted, gene clusters were not resampled, and the intervals do not represent historical model-selection uncertainty, source-label circularity or transport to unresolved variants. They supplement the five-repeat split-stability SDs rather than reinterpret them as confidence intervals.',body)
pp(d,'Table S2. Conditional first-repeat intervals on 2,046 evaluable anchors.',cap)
z=pd.read_csv(ROOT/'tables/conditional_intervals_repeat1.csv');rows=[['Metric','First-repeat estimate','Conditional 2.5th–97.5th percentiles']]+[[r.metric,f'{r.estimate:.4f}',f'{r.conditional_lower:.4f}–{r.conditional_upper:.4f}'] for _,r in z.iterrows()]
addtable(rows,[3400,2700,4360])
pp(d,'S3. gnomAD candidate eligibility',head)
pp(d,'This is an eligibility audit of a selected common-variant stratum. All 50 screened candidates and their frequency, denominator, transcript and overlap records are provided in audit/gnomAD_candidates.csv. Forty-seven overlap the complete source training archive and one additional candidate overlaps the unresolved archive. Both remaining candidates are missense on alternative transcripts, but not on MANE Select. No model performance is calculated from this set.',body)
p=d.add_paragraph();p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(ROOT/'figures/FigureS1_gnomAD_eligibility.png'),width=Inches(7.0))
pp(d,'Figure S1. (A) Candidates by gene and archive overlap. (B) Sequential retention under full-training exclusion, unresolved-archive exclusion and MANE missense restriction. These counts describe the specified gnomAD query and annotation audit; they do not prove that no other candidate could exist under a different release or transcript definition.', 'MDPI_5.1_figure_caption')
pp(d,'Table S3. Two candidates outside both archives; no model predictions.',cap)
rows=[['Gene','GRCh38 variant','Alternative-transcript missense','MANE consequence','Scoring status'],['CHEK2','22-28712124-A-C','ENST00000464581.7:c.11T>G; p.Leu4Arg','Intronic','Unscored'],['PTEN','10-87864144-G-C','ENST00000693560.1:c.194G>C; p.Cys65Ser','5′ UTR','Unscored']]
addtable(rows,[1000,2350,4100,1650,1360])
pp(d,'An alternate-allele frequency screen does not by itself complete a BA1 assessment. Disease and gene specifications, exception handling, appropriate population denominators, transcript relevance and source quality require adjudication. Selection on frequency can also share information with model inputs. The presence of both reference classes in a single database is not a requirement for a descriptive challenge, but mixing class-specific sources does not create representative accuracy or predictive-value estimates.',body)
pp(d,'Methodological source: Ghosh et al. Updated recommendation for the benign stand-alone ACMG/AMP criterion. Hum. Mutat. 2018;39:1525–1530. doi:10.1002/humu.23642. https://clinicalgenome.org/docs/updated-recommendation-for-the-benign-stand-alone-acmg-amp-criterion/',body)
pp(d,'S4. Source provenance and reproducibility',head)
pp(d,'The original 25-feature matrix is retained in data/legacy_reference/ for provenance only; its missing original fold and fitted-model artifacts cannot be reconstructed by asserting that new fits are the old ones. Current results use anchors_all_numeric.csv and the R4 fitted objects. The deployment panel is provided separately as anchor_deployment_panel.csv. Each assessment output includes SPDI, gene, source label, learner probabilities and votes, V-state, Mean4, S_MAC, completeness and partition ID. Inner fold memberships and training identities are recorded in each fit audit. All 58 saved fitted objects and the full application catalogue passed prediction replay to numerical precision (maximum absolute probability difference 1.11 × 10⁻¹⁶).',body)
pp(d,'The historical training sheet lacks original transcript/consequence alignment. Current Ensembl annotation identifies 2,039 MANE missense anchors, 81 additional same-gene alternative-transcript missense anchors and 40 without a current same-gene missense consequence. The 2,120-record sensitivity refit is reported separately. The complete hospital audit retains all 122 source rows, including 75 outside the gene scope. The 47 in-scope rows resolve to 31 missense variants, all scored in Table 7: 22 anchor overlaps, eight application-catalogue overlaps and one non-anchor source-archive record. All meet ≥20/25 completeness. Anchor-overlapping deployment scores are in-sample. Repeated rows are not assumed to be independent patients; clinical reference labels remain unadjudicated.',body)
pp(d,'Transcript identity was checked against Ensembl snapshots and NCBI ClinVar VCV000005591.136 (https://www.ncbi.nlm.nih.gov/clinvar/variation/5591/; accessed 6 September 2026). CHEK2 Ile200Thr/Ile157Thr and Gly210Arg/Gly167Arg were each merged by genomic allele. A bare c.599T>C catalogue lookup would select p.Val200Ala on MANE Select. The full row audit preserves source names, transcript strings and discrepant ATM rsIDs. Additional Ensembl responses are in audit/cache/hospital_additional_transcripts.json.gz. These checks establish allele identity, not clinical reference classifications.',body)
pp(d,'S5. TRIPOD+AI reporting map and remaining gaps',head)
pp(d,'This is a study-specific reporting map, not a claim of complete compliance. Item topics are abbreviated and the official checklist remains authoritative. Patient-level prediction items are interpreted cautiously because the analysis unit is a variant. Persistent deposition, historical annotation-release provenance and some author declarations remain incomplete; a supplementary checklist cannot remedy these gaps.',body)
MAP=[
('1–2','Title and abstract','Title; Abstract','Model, target, outputs, evaluability and internal scope stated.'),
('3a–c','Context and users','Introduction; Discussion','Research-curator use specified; ancestry/demographic representativeness unavailable.'),
('4','Objectives','Introduction final paragraph; 4.1','Development and internal evaluation; clinical validation withdrawn.'),
('5a–b','Source data and dates','4.1–4.3; input_manifest.json; audit/cache','Source archives hashed; 2026-09-05 audit snapshots retained. Historical source release dates unavailable.'),
('6a–c','Eligibility and setting','4.1–4.2; Table 1','Variant-level scope; historical transcript discrepancies declared; treatment/follow-up not applicable.'),
('7','Preparation','4.3; prepare_inputs.py','Numeric parsing, semantic reduction and partition-specific preprocessing documented.'),
('8a–c','Reference labels','4.2; Discussion; source metadata','Processed y distinguished from embedded ClinVar labels. Independent blinded outcome adjudication unavailable.'),
('9a–c','Predictors','4.3; Table 3; data dictionary','47 primary candidates and 25 deployment predictors supplied. Upstream training provenance remains incomplete.'),
('10','Sample size','4.1; Tables 1, 5–6','All available source anchors used; no formal prospective calculation; small gene strata acknowledged.'),
('11','Missingness','2.2; 4.2–4.3; output flags','Median imputation and identical completeness suppression; all denominators retained.'),
('12a–g','Analysis design','4.4–4.9; scripts; fold maps','Nested fitting, tuning, calibration, thresholding, metrics and sensitivity analyses explicit.'),
('13','Class balance','4.1; fitted learner settings','No class resampling or class-weight adjustment; class counts and calibration reported.'),
('14','Fairness','Discussion; data dictionary','Ancestry-specific and demographic fairness not estimable from these source data.'),
('15','Outputs','4.2 and 4.6; Table 5','Shared vote descriptors; retrospective convention; S_MAC ranking; completeness guard.'),
('16','Population differences','2.7; 4.8; Figure 7','Predictor shift, disagreement shift, gene information and absent application outcomes documented.'),
('17','Ethics','Ethics and consent statements','Source hospital approval retained; no new external clinical outcome analysis.'),
('18a–b','Funding and interests','Acknowledgments; conflicts statement','Conflict declaration retained. Funding declaration requires author completion.'),
('18c–d','Protocol and registration','4.1; executed_protocol.md','Revision-stage reanalysis not prospectively registered; executed protocol and deviations supplied.'),
('18e–f','Data and code access','Availability statements; release archive','Matrices, folds, outputs, models and scripts included. Persistent DOI and reviewer access still pending.'),
('19','Public involvement','Author declaration needed','No documented patient/public involvement record was supplied; authors must confirm the statement.'),
('20a–c','Cohort accounting','Tables 1 and 5; Figures 1, 3, 7','Source and evaluable counts, application population and missingness distinguished; clinical demographics unavailable.'),
('21','Analysis denominators','Tables 5–6; per-gene tables; fit audits','Training and held-out identities, class counts and incomplete flags provided per analysis.'),
('22','Complete model','Table 4; models/deployment.joblib','Preprocessors, parameters, calibrators, thresholds and scoring code supplied; persistent access pending.'),
('23a–b','Performance and variation','2.2–2.7; Figures 2–7; Tables S1–S2','Repeat SDs, per-gene Wilson intervals and conditional fixed-score intervals. No prospective or transport intervals.'),
('24','Model revision','2.1; README; executed_protocol.md','New fits supersede prior numerical estimates; old missing fits are not reconstructed or claimed verified.'),
('25–26','Interpretation and limits','Discussion','Source-label dependence, gene confounding, shift, cost asymmetry, missingness and transcript limits connected.'),
('27a–c','Use and next evaluation','Discussion; 4.2; viewer notes','Expert research use; verify build/transcript; incomplete suppression; independent evidence and clinical utility evaluation remain necessary.')]
pp(d,'Table S4. Study-specific TRIPOD+AI reporting map.',cap)
rows=[['Item','Topic','Location','Coverage / limitation']]+[list(x) for x in MAP]
reporting_table=addtable(rows,[700,1800,2750,5210])
for spacing in reporting_table.findall('.//w:pPr/w:spacing',NS):
 spacing.set(qn('w:before'),'0');spacing.set(qn('w:after'),'0')
for size in reporting_table.findall('.//w:rPr/w:sz',NS):size.set(qn('w:val'),'16')
pp(d,'Checklist: Collins et al. TRIPOD+AI. BMJ 2024;385:e078378. doi:10.1136/bmj-2023-078378. Expanded checklist: https://www.tripod-statement.org/wp-content/uploads/2024/04/TRIPODAI-Supplement.pdf.',body)
# Keep headings and captions with following content; table rows may split only between rows.
for p in d.paragraphs:
 if p.style.style_id in ['MDPI21heading1','MDPI22heading2','MDPI41tablecaption']:p.paragraph_format.keep_with_next=True
 if p.style.style_id=='MDPI51figurecaption':p.paragraph_format.keep_together=True
for t in d.tables:
 for row in t.rows:
  trpr=row._tr.get_or_add_trPr()
  if trpr.find(qn('w:cantSplit')) is None:trpr.append(OxmlElement('w:cantSplit'))
save_document(d,OUT/'HBOCpred_Supplementary_Audits_TRIPODAI.docx')
(ROOT/'documentation/TRIPODAI_reporting_map.csv').write_text(pd.DataFrame(MAP,columns=['item','topic','location','coverage_and_gap']).to_csv(index=False))
print('Response and Supplementary Materials written')
